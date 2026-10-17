"""Gera a planilha de fluxo de caixa do Motoca Food.

Uso: python3 gerar_fluxo.py <saida.xlsx> <AAAA-MM do 1º mês>
Os valores em PREMISSAS são o ponto de partida; edite-os aqui ou direto na planilha (células amarelas).
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SALDO_INICIAL = 0  # não informado: substitua pelo saldo da conta PJ
RESERVA_MESES = 2
# (categoria, tipo, valor mensal previsto, dia, observação)
PREMISSAS = [
    ("Patrocínio", "Entrada", 0, "", "Só entra com contrato assinado (Parceiro Fundador é sem custo)"),
    ("Outras receitas", "Entrada", 0, "", ""),
    ("Impostos (DAS-MEI)", "Saída", 86.05, 20, "Valor real 2026, serviços: 5% de R$1.621 + R$5 ISS. Vence dia 20"),
    ("App/tecnologia", "Saída", 300, "", "ESTIMATIVA: substitua (servidor, APIs, lojas de app)"),
    ("Marketing", "Saída", 500, "", "ESTIMATIVA: substitua (impulsionamento, vídeos)"),
    ("Ferramentas/assinaturas", "Saída", 250, "", "ESTIMATIVA: substitua (IA, design, telefone)"),
    ("Pessoas", "Saída", 800, "", "ESTIMATIVA: substitua (freelancer)"),
    ("Tarifas bancárias", "Saída", 0, "", "Confira no extrato da conta PJ"),
    ("Imprevistos", "Saída", 100, "", "ESTIMATIVA: colchão para gastos não previstos"),
]
CATEGORIAS = ["Aporte do titular"] + [p[0] for p in PREMISSAS[:2]] + [p[0] for p in PREMISSAS[2:]] + ["Anuais/pontuais"]

ARIAL = "Arial"
F = lambda **k: Font(name=ARIAL, **k)
BLUE, GREEN = "0000FF", "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F3A5F")
SECTION = PatternFill("solid", fgColor="DCE6F1")
MONEY = 'R$ #,##0.00;[Red](R$ #,##0.00);"-"'
THIN = Border(bottom=Side(style="thin", color="999999"))


def header(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build(path, start):
    wb = Workbook()

    # ---------- Premissas ----------
    pr = wb.active
    pr.title = "Premissas"
    pr["A1"] = "Premissas do fluxo de caixa: Motoca Food (MEI)"
    pr["A1"].font = F(bold=True, size=14)
    pr["A2"] = "Edite só as células AMARELAS (texto azul). O resto é calculado."
    pr["A2"].font = F(italic=True, color="555555")
    for r, (label, val, fmt, note) in enumerate([
        ("Saldo inicial da conta PJ (R$)", SALDO_INICIAL, MONEY, "Não informado: coloque o saldo de hoje do extrato PJ"),
        ("Reserva mínima (meses de saídas)", RESERVA_MESES, "0", "Quantos meses de gasto devem sobrar sempre na conta"),
        ("Primeiro mês do fluxo", start, "mmm/yyyy", ""),
    ], start=3):
        pr.cell(row=r, column=1, value=label).font = F(bold=True)
        c = pr.cell(row=r, column=2, value=val)
        c.font, c.fill, c.number_format = F(color=BLUE), YELLOW, fmt
        pr.cell(row=r, column=3, value=note).font = F(italic=True, color="555555")
    header(pr, 7, ["Categoria", "Tipo", "Valor mensal previsto (R$)", "Dia", "Observação / fonte"])
    for i, (cat, tipo, val, dia, obs) in enumerate(PREMISSAS, start=8):
        pr.cell(row=i, column=1, value=cat).font = F()
        pr.cell(row=i, column=2, value=tipo).font = F()
        c = pr.cell(row=i, column=3, value=val)
        c.font, c.fill, c.number_format = F(color=BLUE), YELLOW, MONEY
        d = pr.cell(row=i, column=4, value=dia)
        d.font, d.fill = F(color=BLUE), YELLOW
        pr.cell(row=i, column=5, value=obs).font = F(italic=True, color="555555")
    last = 7 + len(PREMISSAS)
    pr.cell(row=last + 2, column=1, value="Aporte do titular e gastos anuais/pontuais são lançados mês a mês na aba 'Fluxo 12 meses' (linhas amarelas).").font = F(italic=True)
    pr["E10"].comment = Comment("Fonte: O Povo, 02/01/2026; Contábeis. Salário mínimo 2026 = R$1.621. Confira no Portal do Empreendedor.", "Motoca Food")
    for col, w in zip("ABCDE", (34, 10, 26, 6, 70)):
        pr.column_dimensions[col].width = w

    # ---------- Lançamentos ----------
    la = wb.create_sheet("Lançamentos")
    header(la, 1, ["Data", "Tipo", "Categoria", "Descrição", "Valor (R$)", "Forma", "Comprovante"])
    ex = [date(2026, 11, 20), "Saída", "Impostos (DAS-MEI)", "EXEMPLO: DAS competência out/2026 (apague esta linha)", 86.05, "Pix", "sim"]
    for j, v in enumerate(ex, 1):
        c = la.cell(row=2, column=j, value=v)
        c.font = F(italic=True, color="808080")
    la["A2"].number_format = "dd/mm/yyyy"
    la["E2"].number_format = MONEY
    for r in range(3, 501):
        la.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        la.cell(row=r, column=5).number_format = MONEY
        for col in range(1, 8):
            la.cell(row=r, column=col).font = F(color=BLUE)
    dv_tipo = DataValidation(type="list", formula1='"Entrada,Saída"', allow_blank=True)
    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORIAS) + '"', allow_blank=True)
    la.add_data_validation(dv_tipo)
    la.add_data_validation(dv_cat)
    dv_tipo.add("B2:B500")
    dv_cat.add("C2:C500")
    for col, w in zip("ABCDEFG", (12, 10, 26, 52, 14, 10, 12)):
        la.column_dimensions[col].width = w
    la.freeze_panes = "A2"

    # ---------- Fluxo 12 meses ----------
    fl = wb.create_sheet("Fluxo 12 meses")
    fl["A1"] = "Fluxo de caixa: previsto 12 meses (R$)"
    fl["A1"].font = F(bold=True, size=14)
    cols = [get_column_letter(c) for c in range(3, 15)]  # C..N
    tot = "O"
    fl["B3"] = "Mês"
    for k, cl in enumerate(cols):
        c = fl[f"{cl}3"]
        c.value = "=Premissas!$B$5" if k == 0 else f"=EDATE({cols[k-1]}3,1)"
        c.number_format, c.font, c.fill = "mmm/yy", F(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(horizontal="center")
    fl[f"{tot}3"] = "Total 12m"
    fl[f"{tot}3"].font, fl[f"{tot}3"].fill = F(bold=True, color="FFFFFF"), HEAD
    fl["B3"].font, fl["B3"].fill = F(bold=True, color="FFFFFF"), HEAD

    rows = {}

    def label(r, text, bold=False, fill=None):
        c = fl.cell(row=r, column=2, value=text)
        c.font = F(bold=bold)
        if fill:
            for col in range(2, 16):
                fl.cell(row=r, column=col).fill = fill

    label(4, "Saldo inicial", bold=True)
    label(5, "ENTRADAS", bold=True, fill=SECTION)
    r = 6
    rows["aporte"] = r
    label(r, "Aporte do titular (dinheiro próprio)")
    for cl in cols:
        c = fl[f"{cl}{r}"]
        c.value, c.font, c.fill = 0, F(color=BLUE), YELLOW
    r += 1
    first_in = r
    for i in range(2):  # Patrocínio, Outras receitas
        label(r, f"=Premissas!A{8+i}")
        fl.cell(row=r, column=2).font = F(color=GREEN)
        for cl in cols:
            fl[f"{cl}{r}"] = f"=Premissas!$C${8+i}"
            fl[f"{cl}{r}"].font = F(color=GREEN)
        r += 1
    rows["sem_aporte"] = (first_in, r - 1)
    rows["tot_in"] = r
    label(r, "Total entradas", bold=True)
    for cl in cols:
        fl[f"{cl}{r}"] = f"=SUM({cl}{rows['aporte']}:{cl}{r-1})"
        fl[f"{cl}{r}"].font = F(bold=True)
    r += 1
    label(r, "SAÍDAS", bold=True, fill=SECTION)
    r += 1
    first_out = r
    for i in range(2, len(PREMISSAS)):
        label(r, f"=Premissas!A{8+i}")
        fl.cell(row=r, column=2).font = F(color=GREEN)
        for cl in cols:
            fl[f"{cl}{r}"] = f"=Premissas!$C${8+i}"
            fl[f"{cl}{r}"].font = F(color=GREEN)
        r += 1
    label(r, "Anuais/pontuais (domínio, lojas de app…)")
    for cl in cols:
        c = fl[f"{cl}{r}"]
        c.value, c.font, c.fill = 0, F(color=BLUE), YELLOW
    r += 1
    rows["tot_out"] = r
    label(r, "Total saídas", bold=True)
    for cl in cols:
        fl[f"{cl}{r}"] = f"=SUM({cl}{first_out}:{cl}{r-1})"
        fl[f"{cl}{r}"].font = F(bold=True)
    r += 1
    rows["res"] = r
    label(r, "Resultado do mês", bold=True)
    for cl in cols:
        fl[f"{cl}{r}"] = f"={cl}{rows['tot_in']}-{cl}{rows['tot_out']}"
    r += 1
    rows["saldo"] = r
    label(r, "Saldo final", bold=True)
    for cl in cols:
        fl[f"{cl}{r}"] = f"={cl}4+{cl}{rows['res']}"
        fl[f"{cl}{r}"].font = F(bold=True)
        fl[f"{cl}{r}"].border = THIN
    r += 1
    rows["reserva"] = r
    label(r, "Reserva mínima")
    for cl in cols:
        fl[f"{cl}{r}"] = f"=Premissas!$B$4*{cl}{rows['tot_out']}"
    r += 1
    rows["falta"] = r
    label(r, "Falta para a reserva")
    for cl in cols:
        fl[f"{cl}{r}"] = f"=MAX(0,{cl}{rows['reserva']}-{cl}{rows['saldo']})"
    r += 1
    rows["flag"] = r
    label(r, "Abaixo da reserva? (1 = sim)")
    for cl in cols:
        fl[f"{cl}{r}"] = f"=IF({cl}{rows['saldo']}<{cl}{rows['reserva']},1,0)"
        fl[f"{cl}{r}"].number_format = "0"
    r += 1
    rows["alerta"] = r
    label(r, "Alerta", bold=True)
    for cl in cols:
        fl[f"{cl}{r}"] = f'=IF({cl}{rows["saldo"]}<0,"NEGATIVO",IF({cl}{rows["flag"]}=1,"abaixo da reserva","ok"))'
        fl[f"{cl}{r}"].alignment = Alignment(horizontal="center")
    # saldo inicial
    for k, cl in enumerate(cols):
        fl[f"{cl}4"] = "=Premissas!$B$3" if k == 0 else f"={cols[k-1]}{rows['saldo']}"
        fl[f"{cl}4"].font = F(color=GREEN if k == 0 else "000000", bold=True)
    r += 2
    label(r, "REALIZADO (puxado da aba Lançamentos)", bold=True, fill=SECTION)
    r += 1
    rows["real_in"], rows["real_out"], rows["desvio"] = r, r + 1, r + 2
    for rr, txt, tipo in [(r, "Entradas realizadas", "Entrada"), (r + 1, "Saídas realizadas", "Saída")]:
        label(rr, txt)
        for cl in cols:
            fl[f"{cl}{rr}"] = (f"=SUMIFS(Lançamentos!$E:$E,Lançamentos!$B:$B,\"{tipo}\","
                               f"Lançamentos!$A:$A,\">=\"&{cl}$3,Lançamentos!$A:$A,\"<\"&EDATE({cl}$3,1))")
            fl[f"{cl}{rr}"].font = F(color=GREEN)
    label(r + 2, "Desvio saídas (realizado − previsto)")
    for cl in cols:
        fl[f"{cl}{r+2}"] = f"={cl}{rows['real_out']}-{cl}{rows['tot_out']}"
    last_row = r + 2
    # totais e formatos
    for rr in range(4, last_row + 1):
        for cl in cols + [tot]:
            if fl[f"{cl}{rr}"].number_format == "General":
                fl[f"{cl}{rr}"].number_format = MONEY
            if fl[f"{cl}{rr}"].font.name != ARIAL:
                fl[f"{cl}{rr}"].font = F()
    for rr in list(range(rows["aporte"], rows["res"] + 1)) + [rows["real_in"], rows["real_out"], rows["desvio"]]:
        if fl[f"B{rr}"].value in ("SAÍDAS",):
            continue
        fl[f"{tot}{rr}"] = f"=SUM(C{rr}:N{rr})"
        fl[f"{tot}{rr}"].font = F(bold=True)
        fl[f"{tot}{rr}"].number_format = MONEY
    fl.column_dimensions["A"].width = 2
    fl.column_dimensions["B"].width = 40
    for cl in cols + [tot]:
        fl.column_dimensions[cl].width = 13
    fl.freeze_panes = "C4"

    # ---------- Resumo ----------
    rs = wb.create_sheet("Resumo", 0)
    rs["A1"] = "Resumo financeiro: Motoca Food"
    rs["A1"].font = F(bold=True, size=14)
    rs["A2"] = "Tudo aqui é calculado. Para mudar, edite as células amarelas em Premissas e no Fluxo."
    rs["A2"].font = F(italic=True, color="555555")
    FL = "'Fluxo 12 meses'"
    a, b = rows["sem_aporte"]
    items = [
        ("Saldo inicial da conta PJ", "=Premissas!B3", MONEY),
        ("Total de entradas em 12 meses", f"={FL}!O{rows['tot_in']}", MONEY),
        ("  dos quais aporte do titular", f"={FL}!O{rows['aporte']}", MONEY),
        ("Total de saídas em 12 meses", f"={FL}!O{rows['tot_out']}", MONEY),
        ("Saldo final do período", f"={FL}!N{rows['saldo']}", MONEY),
        ("Queima média por mês (saídas − receitas, sem aporte)", f"=({FL}!O{rows['tot_out']}-SUM({FL}!O{a}:O{b}))/12", MONEY),
        ("Runway: meses que o saldo inicial cobre", "=IF(B8<=0,\"sem queima\",B3/B8)", "0.0"),
        ("Menor saldo final no período", f"=MIN({FL}!C{rows['saldo']}:N{rows['saldo']})", MONEY),
        ("Aporte extra necessário para nunca ficar abaixo da reserva", f"=MAX(0,MAX({FL}!C{rows['falta']}:N{rows['falta']}))", MONEY),
        ("  em média por mês", "=B11/12", MONEY),
        ("Primeiro mês abaixo da reserva", f"=IFERROR(INDEX({FL}!C3:N3,MATCH(1,{FL}!C{rows['flag']}:N{rows['flag']},0)),\"nenhum\")", "mmm/yyyy"),
    ]
    for i, (lab, frm, fmt) in enumerate(items, start=3):
        rs.cell(row=i, column=1, value=lab).font = F(bold=not lab.startswith("  "))
        c = rs.cell(row=i, column=2, value=frm)
        c.font, c.number_format = F(bold=True), fmt
    rs["A15"] = "Leitura: com o saldo inicial e a queima média, o dinheiro dura B9 meses. Para manter a reserva em todos os meses, aporte B11 ao longo do período."
    rs["A15"].font = F(italic=True)
    rs["A16"] = "MEI: DAS vence dia 20; limite de faturamento de R$81 mil/ano (aporte não conta); DASN-SIMEI até 31/05."
    rs["A16"].font = F(italic=True)
    rs.column_dimensions["A"].width = 58
    rs.column_dimensions["B"].width = 18

    wb.save(path)


if __name__ == "__main__":
    out, ym = sys.argv[1], sys.argv[2]
    y, m = map(int, ym.split("-"))
    build(out, date(y, m, 1))
